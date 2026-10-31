"""DOCX export helpers for generated cover letters."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from services.cover_letter_format import (
    candidate_address_lines,
    company_address_lines,
    cover_letter_body_paragraphs,
    cover_letter_date,
    cover_letter_salutation,
    cover_letter_signoff,
    cover_letter_subject,
)


def _safe_line(value: Any) -> str:
    return " ".join(str(value or "").split()).strip()


def build_cover_letter_filename(job: dict[str, Any]) -> str:
    title = "".join(char if char.isalnum() else "_" for char in _safe_line(job.get("title") or "cover_letter"))
    company = "".join(char if char.isalnum() else "_" for char in _safe_line(job.get("company") or "company"))
    title = title.strip("_") or "cover_letter"
    company = company.strip("_") or "company"
    return f"{title[:40]}_{company[:30]}_cover_letter.docx"


def export_cover_letter_docx(profile: dict[str, Any], job: dict[str, Any], package: dict[str, Any]) -> tuple[str, bytes]:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError("DOCX export is unavailable because python-docx is not installed.") from exc

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    for line in candidate_address_lines(profile):
        if line:
            document.add_paragraph(line)
    document.add_paragraph(cover_letter_date())
    document.add_paragraph()

    for line in company_address_lines(job):
        if line:
            document.add_paragraph(line)
    document.add_paragraph()

    subject = document.add_paragraph()
    subject.add_run(cover_letter_subject(job)).bold = True
    document.add_paragraph()

    document.add_paragraph(cover_letter_salutation(job))
    document.add_paragraph()

    for paragraph_text in cover_letter_body_paragraphs(package):
        document.add_paragraph(paragraph_text)

    document.add_paragraph()
    document.add_paragraph(cover_letter_signoff(job))
    if _safe_line(profile.get("name")):
        document.add_paragraph(_safe_line(profile.get("name")))

    buffer = BytesIO()
    document.save(buffer)
    return build_cover_letter_filename(job), buffer.getvalue()
